import pytest
from docx import Document

from app.services import ingestion
from app.services.ingestion import chunk_pages, clean_text, parse_document


def test_clean_text_collapses_whitespace_and_removes_null_bytes():
    messy = "Hello    world\x00\n\n\n\n\nNext   paragraph  "
    assert clean_text(messy) == "Hello world\n\nNext paragraph"


def test_parse_txt_returns_one_page(tmp_path):
    file = tmp_path / "notes.txt"
    file.write_text("First line.\n\nSecond   paragraph.", encoding="utf-8")
    assert parse_document(file) == [{"page": 1, "text": "First line.\n\nSecond paragraph."}]


def test_parse_docx_reads_paragraphs(tmp_path):
    file = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Alpha paragraph.")
    doc.add_paragraph("Beta paragraph.")
    doc.save(str(file))

    pages = parse_document(file)

    assert len(pages) == 1
    assert pages[0]["page"] == 1
    assert "Alpha paragraph." in pages[0]["text"]
    assert "Beta paragraph." in pages[0]["text"]


class FakePage:
    def __init__(self, text):
        self._text = text

    def extract_text(self):
        return self._text


class FakeReader:
    """Pretends to be a 3-page PDF whose middle page has no text."""

    def __init__(self, path):
        self.pages = [FakePage("Page one text"), FakePage(""), FakePage("Page three  text")]


def test_parse_pdf_keeps_real_page_numbers_and_skips_empty_pages(monkeypatch, tmp_path):
    monkeypatch.setattr(ingestion, "PdfReader", FakeReader)

    pages = parse_document(tmp_path / "fake.pdf")

    assert pages == [
        {"page": 1, "text": "Page one text"},
        {"page": 3, "text": "Page three text"},
    ]


def test_parse_document_rejects_unsupported_file_type(tmp_path):
    with pytest.raises(ValueError):
        parse_document(tmp_path / "image.png")


def test_chunks_keep_source_page_and_size_limit(monkeypatch):
    monkeypatch.setattr(ingestion, "CHUNK_SIZE", 100)
    monkeypatch.setattr(ingestion, "CHUNK_OVERLAP", 20)
    pages = [{"page": 1, "text": "word " * 100}, {"page": 2, "text": "other " * 100}]

    chunks = chunk_pages(pages, source="book.pdf")

    assert len(chunks) > 2
    assert all(chunk["source"] == "book.pdf" for chunk in chunks)
    assert {chunk["page"] for chunk in chunks} == {1, 2}
    assert [chunk["chunk_id"] for chunk in chunks] == list(range(len(chunks)))
    assert all(len(chunk["text"]) <= 100 for chunk in chunks)


def test_neighbouring_chunks_overlap(monkeypatch):
    monkeypatch.setattr(ingestion, "CHUNK_SIZE", 60)
    monkeypatch.setattr(ingestion, "CHUNK_OVERLAP", 20)
    words = [f"w{i:02d}" for i in range(40)]

    chunks = chunk_pages([{"page": 1, "text": " ".join(words)}], source="x.txt")

    last_word_of_first_chunk = chunks[0]["text"].split()[-1]
    assert last_word_of_first_chunk in chunks[1]["text"].split()